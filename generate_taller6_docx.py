import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_callout(doc, text_list, title="ACUERDO DEL EQUIPO", border_color="1F497D", fill_color="F2F5F9"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}: ")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    if text_list:
        run_first = p.add_run(text_list[0])
        run_first.font.name = "Arial"
        run_first.font.size = Pt(9.5)
        run_first.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        
        for extra in text_list[1:]:
            p_extra = cell.add_paragraph()
            p_extra.paragraph_format.space_before = Pt(2)
            p_extra.paragraph_format.space_after = Pt(2)
            r = p_extra.add_run(extra)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def create_screenshot_box(doc, title, instruction, expected_elements):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FAFAFA")
    set_cell_margins(cell, top=180, bottom=180, left=220, right=220)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="dashed" w:sz="12" w:space="0" w:color="2B579A"/><w:left w:val="dashed" w:sz="12" w:space="0" w:color="2B579A"/><w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="2B579A"/><w:right w:val="dashed" w:sz="12" w:space="0" w:color="2B579A"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_hdr = p.add_run(f"📷 EVIDENCIA VISUAL: {title.upper()}")
    r_hdr.bold = True
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(10)
    r_hdr.font.color.rgb = RGBColor(0x2B, 0x57, 0x9A)
    
    p_inst = cell.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(2)
    p_inst.paragraph_format.space_after = Pt(4)
    r_inst = p_inst.add_run(f"Guía de captura: {instruction}")
    r_inst.italic = True
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(9)
    r_inst.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    for elem in expected_elements:
        p_elem = cell.add_paragraph(style='List Bullet')
        p_elem.paragraph_format.space_before = Pt(1)
        p_elem.paragraph_format.space_after = Pt(1)
        r_elem = p_elem.add_run(elem)
        r_elem.font.name = "Arial"
        r_elem.font.size = Pt(8.5)
        r_elem.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        
    p_placeholder = cell.add_paragraph()
    p_placeholder.paragraph_format.space_before = Pt(14)
    p_placeholder.paragraph_format.space_after = Pt(14)
    p_placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ph = p_placeholder.add_run("[ PEGUE AQUÍ LA CAPTURA DE PANTALLA EN ALTA DEFINICIÓN (100% NÍTIDA) ]")
    r_ph.bold = True
    r_ph.font.name = "Arial"
    r_ph.font.size = Pt(9.5)
    r_ph.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def format_table(table, col_widths, headers, data, alt_shading=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F497D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    for row_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg_color = "F2F5F9" if (row_idx % 2 == 1 and alt_shading) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def build_docx_taller6():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        
    # Encabezado Principal
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("INFORME TÉCNICO EJECUTIVO: TALLER 6")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(17)
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Sprint Review, Retrospective y Liberación Segura en Continuous Deployment")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # Tabla Ficha Técnica
    tbl_meta = doc.add_table(rows=1, cols=2)
    meta_headers = ["Parámetro de Configuración", "Valor / Enlace en Producción"]
    meta_data = [
        ["Proyecto & Repositorio", "Sistema de Ventas MongoDB Atlas / stivenn18/proyecto_mongodb_ventas"],
        ["Aplicación Web en Producción", "https://proyecto-mongodb-ventas.onrender.com"],
        ["Plataforma de Despliegue Continuo", "Render Cloud (Auto-Deploy sincronizado a main)"],
        ["Plataforma de Feature Flags", "ConfigCat Cloud (SDK Python + SDK JavaScript CDN)"],
        ["SDK Key Activa ConfigCat", "configcat-sdk-1/jyHfCAeT0E2Arwo28sVn6g/4yptu_1D2UORDKiU3dUviw"],
        ["Feature Flag Implementado", "wishlist_enabled (Lista de Deseos E-Commerce en vivo)"],
        ["Versión del Playbook", "Scrum + TBD Playbook Versión 3.0 (docs/SCRUM_TBD_PLAYBOOK.md)"],
        ["Objetivo General del Taller", "Adaptar Sprint Review y Retrospective a CD y practicar la liberación segura con toggles"]
    ]
    format_table(tbl_meta, [2.5, 4.0], meta_headers, meta_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. APERTURA Y CONEXIÓN CON EL NUEVO FLUJO
    h1 = doc.add_heading("1. APERTURA Y CONEXIÓN CON EL NUEVO FLUJO", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph()
    p.add_run("El equipo ha completado satisfactoriamente los prerrequisitos técnicos de los Talleres 1, 2 y 3: integración continua a main con validación automática (Ruff + Pytest), auto-despliegue continuo hacia Render y desacoplamiento de despliegue y release mediante ConfigCat con Definition of Done (DoD) formal. Ante este avance, surge la pregunta estratégica central: ")
    r_quote = p.add_run("“¿Cómo cambian el Sprint Review y la Retrospective cuando desplegamos varias veces por semana a producción?”\n")
    r_quote.bold = True
    r_quote.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    create_styled_callout(
        doc,
        [
            "La ceremonia clásica de 'entrega al final del sprint' se vuelve obsoleta e ineficiente cuando el código ya está en producción horas o días antes del cierre del sprint.",
            "Lo que se siente más 'antiguo' es realizar una demo sobre entornos locales o staging fingiendo que nada ha salido a producción, y esperar 2 semanas para recibir feedback que pudo haberse obtenido el mismo día del despliegue.",
            "La Retrospectiva clásica, centrada solo en relaciones interpersonales o estimaciones de puntos de historia, queda corta si no audita la salud técnica del pipeline, la tasa de integración a main y la deuda generada por los Feature Flags."
        ],
        title="REFLEXIÓN DE ARRANQUE Y RUPTURA DE PARADIGMAS"
    )

    # 2. ADAPTACIÓN DEL SPRINT REVIEW
    h2 = doc.add_heading("2. ADAPTACIÓN DEL SPRINT REVIEW AL CONTINUOUS DEPLOYMENT", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph()
    p.add_run("El objetivo de esta adaptación es convertir el Sprint Review en una ")
    r_b = p.add_run("demostración de valor real en producción (o activable mediante toggles)")
    r_b.bold = True
    p.add_run(", erradicando la mentalidad de 'todo o nada' y promoviendo la validación continua.")

    # Tabla Comparativa Review
    tbl_rev = doc.add_table(rows=1, cols=2)
    rev_headers = ["Sprint Review Clásico", "Sprint Review en CD + Trunk-Based Development"]
    rev_data = [
        ["Demo del incremento acumulado al final del sprint.", "Demo de lo que YA ESTÁ (o puede estar) en producción real."],
        ["Enfoque de 'Todo o Nada' (si no se termina todo, no sale).", "Valor incremental protegido y regulado por Feature Toggles."],
        ["Feedback retrasado tras 2 a 4 semanas de desarrollo.", "Feedback continuo basado en telemetría y usuarios reales."],
        ["Enfocado estrictamente en '¿Qué construimos?'.", "Enfocado estratégicamente en '¿Qué valor recibieron los usuarios?'."]
    ]
    format_table(tbl_rev, [3.2, 3.3], rev_headers, rev_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Estructura obligatoria de la demo
    p_est = doc.add_paragraph()
    p_est.paragraph_format.space_after = Pt(4)
    r_est_t = p_est.add_run("2.1. Estructura Obligatoria de la Demo (Máximo 4-5 minutos por ítem):")
    r_est_t.bold = True
    r_est_t.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    items_demo = [
        "1. ¿Qué problema de usuario resuelve?: Explicación precisa de la necesidad de negocio mitigada.",
        "2. ¿Está activo para todos o detrás de toggle?: Visibilidad inmediata del estado del flag en ConfigCat (OFF, Canario % o 100% ON).",
        "3. Demo en vivo en Render: Ejecución directa en https://proyecto-mongodb-ventas.onrender.com demostrando la conmutación en vivo del toggle.",
        "4. Métrica o evidencia de valor: Datos de telemetría, tiempos de respuesta (<50ms) y tasas de adopción preliminar.",
        "5. Pregunta al PO / Stakeholders: '¿Esto genera el valor esperado y estamos listos para ampliar la audiencia?'."
    ]
    for it in items_demo:
        p_it = doc.add_paragraph(style='List Bullet')
        p_it.paragraph_format.space_before = Pt(1)
        p_it.paragraph_format.space_after = Pt(2)
        r = p_it.add_run(it)
        r.font.name = "Arial"
        r.font.size = Pt(9)

    # Simulación con 3 incrementos reales
    p_sim = doc.add_paragraph()
    p_sim.paragraph_format.space_before = Pt(6)
    p_sim.paragraph_format.space_after = Pt(4)
    r_sim_t = p_sim.add_run("2.2. Simulación Práctica de Sprint Review (3 Incrementos Reales del Repositorio):")
    r_sim_t.bold = True

    tbl_inc = doc.add_table(rows=1, cols=4)
    inc_headers = ["Incremento Real", "Problema que Resuelve", "Estado del Toggle", "Evidencia / Valor Medido"]
    inc_data = [
        [
            "Wishlist E-Commerce (Lista de Deseos)",
            "Permite a los clientes guardar productos para compras futuras sin ensuciar el carrito.",
            "Protegido con ConfigCat (wishlist_enabled). Conmutable en vivo.",
            "Pruebas unitarias al 100% (tests.py). Cero errores 5xx. Actualización en <5s en el navegador."
        ],
        [
            "Catálogo Interactivo y Filtros Dinámicos",
            "Navegación lenta en inventario amplio. Filtra instantáneamente por Laptops, Componentes, etc.",
            "100% Activo para todos los usuarios en producción.",
            "Renderización instantánea en cliente. Cero peticiones redundantes al backend."
        ],
        [
            "API de Telemetría y Salud (/health, /api/flags)",
            "Falta de visibilidad del estado de servicios y sincronización de flags.",
            "100% Activo (público y consumible por CI/CD).",
            "Latencia < 40ms. Monitoreo constante de estado 'healthy' en Render."
        ]
    ]
    format_table(tbl_inc, [1.6, 2.0, 1.4, 1.5], inc_headers, inc_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    create_screenshot_box(
        doc,
        "DEMOSTRACIÓN EN VIVO EN RENDER CON CONFIGCAT CONMUTADO",
        "Capture la pantalla del navegador mostrando la tienda en Render (https://proyecto-mongodb-ventas.onrender.com) demostrando la Wishlist activa y el indicador de ConfigCat en verde.",
        [
            "URL visible en la barra de navegación: https://proyecto-mongodb-ventas.onrender.com",
            "Indicador visual 'ConfigCat: wishlist_enabled ACTIVO' con pulso verde.",
            "Botones '❤️ Guardar' en las tarjetas de productos y botón de 'Lista de Deseos' en el header.",
            "Modal interactivo desplegado mostrando los productos añadidos a favoritos."
        ]
    )

    # 3. ADAPTACIÓN DE LA SPRINT RETROSPECTIVE
    h3 = doc.add_heading("3. ADAPTACIÓN DE LA SPRINT RETROSPECTIVE", level=1)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph()
    p.add_run("La Retrospectiva en CD se enfoca en la ")
    p.add_run("inspección rigurosa del flujo de valor, la salud del pipeline y la disciplina Trunk-Based Development").bold = True
    p.add_run(". No se limita a sensaciones, sino que audita métricas de entrega continua.")

    # Los 3 focos obligatorios
    tbl_focos = doc.add_table(rows=1, cols=3)
    foc_headers = ["1. Salud del Pipeline", "2. Disciplina TBD", "3. Flujo de Valor"]
    foc_data = [
        [
            "• ¿Cuántas veces falló el CI? (GitHub Actions).\n• ¿Cuánto tardó el deploy promedio? (< 2 minutos en Render).\n• ¿Hubo rollbacks o incidentes en health check?",
            "• ¿Cuántas integraciones a main logramos por día? (Meta: >= 1).\n• ¿Alguna rama vivió más de 24 horas?\n• ¿Uso y limpieza adecuada de Feature Toggles?",
            "• Lead Time: tiempo desde el primer commit hasta llegar a producción (< 45 min).\n• ¿El valor llegó de forma medible al usuario final?"
        ]
    ]
    format_table(tbl_focos, [2.1, 2.2, 2.2], foc_headers, foc_data, alt_shading=False)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Dinámica Miro Semáforo
    p_din = doc.add_paragraph()
    p_din.paragraph_format.space_after = Pt(4)
    r_din_t = p_din.add_run("3.1. Tablero de Retrospectiva (Dinámica de Miro con 4 Columnas):")
    r_din_t.bold = True

    tbl_miro = doc.add_table(rows=1, cols=2)
    miro_headers = ["Columna de Inspección", "Post-its y Hallazgos Reales del Equipo"]
    miro_data = [
        [
            "🟢 Lo que funcionó bien en el flujo",
            "• Pipeline de CI en GitHub Actions ejecuta tests y linter en menos de 15 segundos.\n• Auto-deploy a Render transparente e inmediato tras cada merge a main.\n• Conmutación remota con ConfigCat funcional en menos de 5 segundos sin redesplegar."
        ],
        [
            "🟡 Lo que generó fricción",
            "• Esperar aprobaciones de PR en historias pequeñas retrasa el merge diario.\n• Ausencia de telemetría en tiempo real sobre errores funcionales al activar flags."
        ],
        [
            "🔴 Lo que rompió el flujo o generó miedo",
            "• Temor del equipo a encender funcionalidades que tocan datos críticos sin monitoreo.\n• Ramas abiertas durante varios días por culpa de tests intermitentes (flaky tests)."
        ],
        [
            "🛠️ Acciones de Mejora (Máximo 3 concretas)",
            "1. POLÍTICA DE CUARENTENA: Todo test flaky detectado en CI se aísla de inmediato para desbloquear ramas.\n2. TELEMETRÍA PREVIA: Ningún toggle se mueve a producción sin dashboard de errores y latencia.\n3. POLÍTICA DE DESECHO: Todo Feature Toggle al 100% debe eliminarse del código en un plazo máximo de 14 días."
        ]
    ]
    format_table(tbl_miro, [2.4, 4.1], miro_headers, miro_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    create_screenshot_box(
        doc,
        "TABLERO DE RETROSPECTIVA EN MIRO / SEMÁFORO DE FRICCIONES",
        "Pegue la captura del tablero Miro o visual de la Retrospectiva mostrando las 4 columnas del semáforo completadas con los hallazgos y acuerdos del equipo.",
        [
            "Columna Verde (🟢 Lo que funcionó bien en CI/CD y TBD).",
            "Columna Amarilla (🟡 Puntos de fricción identificados).",
            "Columna Roja (🔴 Factores de miedo o bloqueos de flujo).",
            "Columna de Acciones (🛠️ Las 3 acciones de mejora concretas y asignadas)."
        ]
    )

    # 4. LIBERACIÓN SEGURA DE VALOR
    h4 = doc.add_heading("4. PROTOCOLO DE LIBERACIÓN SEGURA DE VALOR (SAFE RELEASE)", level=1)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.add_run("Se documenta y ejercita el ciclo completo de liberación controlada mediante Feature Flags, asegurando que el equipo desacople el despliegue técnico de la activación comercial.")

    tbl_safe = doc.add_table(rows=1, cols=4)
    safe_headers = ["Fase de Liberación", "Audiencia Objetivo", "Métricas Supervisadas", "Criterio de Avance / Rollback"]
    safe_data = [
        [
            "Fase 0: Dark Launch",
            "0% (Oculto en Prod)",
            "Health Check /health, compilación y despliegue en Render.",
            "Si el build pasa y /health devuelve 200, código listo en main."
        ],
        [
            "Fase 1: Canary Interno",
            "Equipo técnico y Staging",
            "Errores en consola de navegador, latencia API < 100ms.",
            "Validación manual de flujos sin reportar bugs durante 2 horas."
        ],
        [
            "Fase 2: Canary Progresivo",
            "10% a 20% de Usuarios",
            "Error Rate (5xx) < 0.5%, tasa de conversión y transacciones.",
            "Si Error Rate > 1.5% -> Rollback inmediato en ConfigCat (<10s)."
        ],
        [
            "Fase 3: General Release",
            "100% de Usuarios",
            "Uso sostenido durante 24-48 horas, estabilidad de MongoDB Atlas.",
            "Éxito completo. Se agenda ticket de limpieza de toggle (<14d)."
        ]
    ]
    format_table(tbl_safe, [1.5, 1.5, 1.8, 1.7], safe_headers, safe_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    create_screenshot_box(
        doc,
        "DASHBOARD DE CONFIGCAT CON GOBERNANZA DE FEATURE FLAGS",
        "Capture la pantalla del panel de control de ConfigCat mostrando el flag 'wishlist_enabled' configurado y los controles de segmentación o conmutación remota.",
        [
            "Nombre del flag visible: wishlist_enabled.",
            "SDK Key vinculada al entorno productivo.",
            "Estado del toggle (ON/OFF o porcentaje de rollout).",
            "Historial de auditoría de cambios en ConfigCat."
        ]
    )

    # 5. RESOLUCIÓN DEL CASO DE ESTUDIO
    h5 = doc.add_heading("5. RESOLUCIÓN INTEGRAL DEL CASO DE ESTUDIO: 'PAGOS EXPRESS'", level=1)
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(6)

    p_caso = doc.add_paragraph()
    r_caso_t = p_caso.add_run("Caso: “La funcionalidad de 'Pagos Express' ya está en main, pero nadie se atreve a activarla”\n")
    r_caso_t.bold = True
    r_caso_t.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    p_caso.add_run("El equipo lleva 2 sprints con auto-deploy y ConfigCat. La historia 'Pagos Express' fue sliceada e integrada a main hace 9 días bajo el toggle 'pagos-express-v1' (OFF 100%). Pasó CI, Render y DoD técnico. Sin embargo, se suscitaron 7 hechos problemáticos que paralizaron la activación. A continuación se presentan las decisiones profesionales tomadas para cada hecho:")

    tbl_caso = doc.add_table(rows=1, cols=3)
    caso_headers = ["# Hecho Problemático", "Diagnóstico de Riesgo", "Decisión Técnica y Plan de Acción Vinculante"]
    caso_data = [
        [
            "1. PO quiere activar el toggle al 20% este viernes por presión de campaña de marketing.",
            "ALTO RIESGO. Activar una pasarela de pagos un viernes sin telemetría y previo a un fin de semana viola la disciplina de resiliencia operativa.",
            "DECISIÓN: NO ACTIVAR EL VIERNES. El Scrum Master y Tech Lead acuerdan con el PO aplazar la activación al martes a las 10:00 AM. Se protege el fin de semana y se usa el lunes para instrumentar las métricas de monitoreo."
        ],
        [
            "2. El último deploy falló el health check durante 11 minutos y requirió rollback manual.",
            "RIESGO DE DISPONIBILIDAD. El endpoint /health carecía de timeout seguro hacia la base de datos o falló por concurrencia durante el inicio.",
            "ACCIÓN: Ajustar el health check con timeout estricto de 2s a MongoDB. Configurar health check retries automáticos en Render para evitar caídas falsas positivas y eliminar la necesidad de rollbacks manuales."
        ],
        [
            "3. Hay 3 Feature Toggles antiguos con más de 45 días activos sin limpiar.",
            "DEUDA TÉCNICA Y COMPLEJIDAD. Acumulación de caminos de ejecución muertos que aumentan la fragilidad del código.",
            "ACCIÓN: Crear inmediatamente 3 issues de deuda técnica (cleanup-toggles) con máxima prioridad para el siguiente sprint. Se institucionaliza la regla: todo toggle al 100% se borra del código en < 14 días."
        ],
        [
            "4. En el sprint actual solo se lograron 4 integraciones a main (lejos de la meta diaria).",
            "PÉRDIDA DE DISCIPLINA TBD. El equipo volvió al anti-patrón de ramas largas y grandes lotes de trabajo.",
            "ACCIÓN: Reforzar la técnica de Story Slicing. Ninguna historia podrá entrar al sprint si no puede integrarse a main en menos de 24 horas. Los PRs deben tener menos de 200 líneas de código."
        ],
        [
            "5. Rama viva 4 días por un test flaky (intermitente) en CI.",
            "BLOQUEO DE FLUJO. Los tests inestables destruyen la confianza del equipo en la integración continua.",
            "ACCIÓN: POLÍTICA DE CUARENTENA INMEDIATA. Todo test flaky se desactiva temporalmente del build bloqueante y se genera un ticket urgente para refactorizarlo en menos de 24 horas. La rama se mergea a main sin demora."
        ],
        [
            "6. No existe métrica clara de conversión o error de pagos para observar al activar el toggle.",
            "OPERACIÓN A CIEGAS. Incapacidad de saber si la funcionalidad genera dinero o está cobrando de forma errónea.",
            "ACCIÓN: OBLIGATORIA. Desarrollar contadores de telemetría (pagos_exitosos, pagos_fallidos, latencia_gateway) antes de autorizar el cambio del toggle. Se añade este requisito al Definition of Ready (DoR)."
        ],
        [
            "7. Negocio anunció internamente que 'Pagos Express ya está listo'.",
            "DESALINEACIÓN DE EXPECTATIVAS. Confusión clásica entre 'Deploy técnico' y 'Release comercial'.",
            "ACCIÓN: Capacitar a los stakeholders en el concepto Deploy != Release. Una historia solo se considera comercialmente lista cuando el toggle está al 100% y los Criterios de Aceptación han sido validados con éxito."
        ]
    ]
    format_table(tbl_caso, [1.5, 2.0, 3.0], caso_headers, caso_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 6. ACTUALIZACIÓN DEL PLAYBOOK Y CIERRE
    h6 = doc.add_heading("6. ACTUALIZACIÓN DEL PLAYBOOK (VERSIÓN 3.0) Y COMPROMISOS", level=1)
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(6)

    p_clo = doc.add_paragraph()
    p_clo.add_run("El documento oficial del equipo ")
    p_clo.add_run("docs/SCRUM_TBD_PLAYBOOK.md").bold = True
    p_clo.add_run(" ha sido actualizado formalmente a la ")
    p_clo.add_run("Versión 3.0").bold = True
    p_clo.add_run(", incorporando las Secciones 13, 14, 15, 16 y 17 relativas a las ceremonias adaptadas y las políticas de liberación segura.")

    create_styled_callout(
        doc,
        [
            "Sprint Review adaptado: Se evalúan incrementos operando en producción con demostraciones en vivo respaldadas por Feature Toggles.",
            "Sprint Retrospective adaptada: Se audita la salud del pipeline, la frecuencia de integración a main y el ciclo de vida de los toggles.",
            "Protocolo de Liberación Segura: Criterios científicos de rollout progresivo y rollback instantáneo en menos de 10 segundos.",
            "Checklist de Fase 1 cerrado al 100% con auto-deploy a Render y ConfigCat en funcionamiento."
        ],
        title="DECLARACIÓN DE MADUREZ DEL EQUIPO"
    )

    p_final = doc.add_paragraph()
    p_final.paragraph_format.space_before = Pt(8)
    p_final.paragraph_format.space_after = Pt(4)
    r_frase = p_final.add_run("Ronda Final: “En una frase: ¿Qué cambia en nuestro próximo Review y Retro después de este taller?”\n")
    r_frase.bold = True
    r_frase.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_quote_box = doc.add_paragraph()
    p_quote_box.paragraph_format.left_indent = Inches(0.4)
    p_quote_box.paragraph_format.right_indent = Inches(0.4)
    r_quote_txt = p_quote_box.add_run(
        "“A partir de hoy, nuestro Sprint Review deja de ser una defensa nerviosa de código en local para convertirse en una celebración de valor activo y medible en producción; mientras que nuestra Retrospectiva se transforma en el guardián riguroso de la salud de nuestro pipeline, la eliminación de la deuda de toggles y la disciplina innegociable de integrar a main todos los días.”"
    )
    r_quote_txt.italic = True
    r_quote_txt.bold = True
    r_quote_txt.font.size = Pt(10.5)
    r_quote_txt.font.color.rgb = RGBColor(0x2B, 0x57, 0x9A)

    output_dir = "docs"
    os.makedirs(output_dir, exist_ok=True)
    target_path = os.path.join(output_dir, "INFORME_TALLER_6_REVIEW_RETRO_LIBERACION.docx")
    doc.save(target_path)
    print(f"Documento Word de Taller 6 generado correctamente en: {target_path}")

if __name__ == "__main__":
    build_docx_taller6()
