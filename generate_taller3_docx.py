import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_callout(doc, text_list, title="ACUERDO DEL EQUIPO", border_color="1F497D", fill_color="F2F5F9"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}: ")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    if text_list:
        run_first = p.add_run(text_list[0])
        run_first.font.name = "Arial"
        run_first.font.size = Pt(9.5)
        run_first.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        
        for extra in text_list[1:]:
            p_extra = cell.add_paragraph()
            p_extra.paragraph_format.space_before = Pt(2)
            p_extra.paragraph_format.space_after = Pt(2)
            r = p_extra.add_run(extra)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_docx_taller3():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("INFORME TÉCNICO: TALLER 3")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("CI/CD, Automatización de Pruebas y Cierre de Pendientes (Render + ConfigCat + DoD Final)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11.5)
    run_sub.bold = True
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Proyecto / Dominio:", "Sistema E-Commerce de Ventas con Wishlist en MongoDB Atlas y FastAPI"),
        ("Repositorio Oficial:", "https://github.com/stivenn18/proyecto_mongodb_ventas"),
        ("Despliegue Render:", "https://proyecto-mongodb-ventas.onrender.com (Auto-Deploy activo desde main)"),
        ("Estado de Fase 1:", "✅ 100% CERRADA Y VALIDADA EN PRODUCCIÓN")
    ]
    col_widths = [Inches(2.0), Inches(4.5)]
    for r_idx, (label, val) in enumerate(meta_data):
        row = meta_tbl.rows[r_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = col_widths[0], col_widths[1]
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
        set_cell_background(cell_lbl, "EBF1F5")
        set_cell_background(cell_val, "F9FAFC")
        
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        r_l = p_lbl.add_run(label)
        r_l.bold = True
        r_l.font.name = "Arial"
        r_l.font.size = Pt(9.5)
        r_l.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        r_v = p_val.add_run(val)
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    h1 = doc.add_heading("1. Diagnóstico del Pipeline Actual (Bloque 1)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    
    p_flow = doc.add_paragraph(
        "Flujo Completo: Push/PR ➔ CI (Pytest + Ruff) ➔ Code Review ➔ Merge a main ➔ Imagen Docker en GHCR ➔ Auto-Deploy a Render ➔ Feature Toggle en ConfigCat"
    )
    p_flow.runs[0].bold = True
    p_flow.paragraph_format.space_after = Pt(8)
    
    sem_tbl = doc.add_table(rows=4, cols=3)
    sem_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["🟢 Ya Funciona Bien", "🟡 Parcial / Manual", "🔴 Falta / Resuelto Hoy"]
    s_widths = [Inches(2.1), Inches(2.1), Inches(2.3)]
    
    for idx, h_t in enumerate(s_headers):
        c = sem_tbl.rows[0].cells[idx]
        c.width = s_widths[idx]
        set_cell_background(c, "1F497D")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p_h = c.paragraphs[0]
        p_h.paragraph_format.space_after = Pt(0)
        r = p_h.add_run(h_t)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    s_rows = [
        ("Pipeline CI en GitHub Actions (<15s) con Ruff y Pytest.", "Configuración manual de variables en ambientes externos.", "Auto-deploy continuo desde main a Render mediante render.yaml."),
        ("Branch Protection y Rulesets en main bloquean código roto.", "Evaluación local de Feature Flags mediante selectores en UI.", "Integración del SDK de ConfigCat para gestión remota de toggles."),
        ("Contenedorización Docker y publicación en GHCR.", "Pruebas manuales de endpoints en Swagger / cliente web.", "Definition of Done (DoD) formalizado en 3 dimensiones (Prod).")
    ]
    for r_idx, s_data in enumerate(s_rows):
        row = sem_tbl.rows[r_idx + 1]
        for cell_idx, val in enumerate(s_data):
            cell = row.cells[cell_idx]
            cell.width = s_widths[cell_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F9FAFC")
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_after = Pt(0)
            r = p_c.add_run(val)
            r.font.size = Pt(8.5)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    h2 = doc.add_heading("2. Automatización de Pruebas y Robustez del CI (Bloque 2)", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    
    ci_rules = [
        "1. Tests Unitarios Obligatorios: Cobertura total de la lógica de WishlistService y endpoints de API.",
        "2. Rutas Estáticas de Frontend: Verificación de disponibilidad de /style.css y /script.js.",
        "3. Linting Estricto: ruff check . ejecutado en modo cero tolerancia para estilo PEP 8.",
        "4. Health Check Automatizado: Test unitario verificando que /health responda status 200 en <100ms.",
        "5. Tiempo Máximo Aceptable del CI: Menos de 60 segundos por ejecución (promedio actual: 12 segundos)."
    ]
    for cr in ci_rules:
        p_cr = doc.add_paragraph(cr, style='List Bullet')
        p_cr.paragraph_format.space_after = Pt(2)
        
    create_styled_callout(
        doc,
        ['"Si el CI está rojo en main o en un PR, se detiene cualquier avance de nuevas funcionalidades hasta que el pipeline vuelva a estar en verde."'],
        title="ACUERDO DE CALIDAD DE EQUIPO"
    )

    h3 = doc.add_heading("3. Auto-Deploy a Render desde main (Bloque 3)", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    
    render_specs = [
        ("URL de Producción:", "https://proyecto-mongodb-ventas.onrender.com"),
        ("Servicio Web:", "ventas-wishlist-app (Ambiente Python 3.11)"),
        ("Rama de Despliegue:", "main (Auto-Deploy activado)"),
        ("Comando de Build:", "pip install -r requirements.txt"),
        ("Comando de Start:", "uvicorn app:app --host 0.0.0.0 --port $PORT"),
        ("Health Check Path:", "/health (Monitoreo de estado y features activas)"),
        ("Variables de Entorno:", "CONFIGCAT_SDK_KEY, MONGO_URI, ENVIRONMENT=production")
    ]
    
    r_tbl = doc.add_table(rows=len(render_specs), cols=2)
    r_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (k, v) in enumerate(render_specs):
        row = r_tbl.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width, c2.width = Inches(2.2), Inches(4.3)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        set_cell_background(c1, "EBF1F5")
        set_cell_background(c2, "FFFFFF")
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(k)
        r1.bold = True
        r1.font.size = Pt(8.5)
        
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(v)
        r2.font.size = Pt(8.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    h4 = doc.add_heading("4. Feature Toggle con ConfigCat (Bloque 4)", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    
    toggle_tbl = doc.add_table(rows=4, cols=3)
    toggle_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Estado del Toggle", "Comportamiento en la Tienda", "Impacto en Usuarios"]
    t_widths = [Inches(1.8), Inches(2.5), Inches(2.2)]
    
    for idx, h_t in enumerate(t_headers):
        c = toggle_tbl.rows[0].cells[idx]
        c.width = t_widths[idx]
        set_cell_background(c, "1F497D")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p_h = c.paragraphs[0]
        p_h.paragraph_format.space_after = Pt(0)
        r = p_h.add_run(h_t)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    t_rows = [
        ("🔴 Toggle OFF\n(0% Rollout)", "Botón '❤️ Guardar' y modal ocultos. Endpoints retornan status disabled.", "Código desplegado en producción pero 100% invisible y seguro."),
        ("🟡 Toggle CANARY\n(10% Rollout)", "Botón visible en tarjetas de producto para grupo beta-testers.", "Validación en vivo con telemetría y feedback real."),
        ("🟢 Toggle ON\n(100% Rollout)", "Funcionalidad de Wishlist accesible para el 100% de clientes.", "Liberación completada ➔ Se programa retiro del código del flag.")
    ]
    for r_idx, t_data in enumerate(t_rows):
        row = toggle_tbl.rows[r_idx + 1]
        for cell_idx, val in enumerate(t_data):
            cell = row.cells[cell_idx]
            cell.width = t_widths[cell_idx]
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F9FAFC")
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_after = Pt(0)
            r = p_c.add_run(val)
            r.font.size = Pt(8.5)
            if cell_idx == 0:
                r.bold = True
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    h5 = doc.add_heading("5. Definition of Done (DoD) Versión Final (Bloque 5)", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    
    dod_tbl = doc.add_table(rows=6, cols=3)
    dod_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    d_headers = ["⚙️ DoD Técnico", "💼 DoD de Negocio", "🚀 DoD de Despliegue"]
    d_widths = [Inches(2.1), Inches(2.2), Inches(2.2)]
    
    for idx, h_t in enumerate(d_headers):
        c = dod_tbl.rows[0].cells[idx]
        c.width = d_widths[idx]
        set_cell_background(c, "1F497D")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p_h = c.paragraphs[0]
        p_h.paragraph_format.space_after = Pt(0)
        r = p_h.add_run(h_t)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    d_items = [
        ("• Código revisado en Pull Request.", "• Acceptance Criteria cumplidos.", "• Auto-desplegado en Render desde main."),
        ("• CI verde (Pytest + Ruff linter).", "• Validable en entorno de destino.", "• Health check /health respondiendo 200 OK."),
        ("• Imagen Docker generada en GHCR.", "• PO conoce el estado del Feature Toggle.", "• Logs y telemetría sin excepciones."),
        ("• Feature Toggle probado en unit tests.", "• Feedback de usuarios registrado.", "• Toggle en estado acordado (OFF/10/100)."),
        ("• No rompe funcionalidad existente.", "• Documentación e issues actualizados.", "• Rollback verificado en caso de falla.")
    ]
    for r_idx, row_data in enumerate(d_items):
        row = dod_tbl.rows[r_idx + 1]
        for cell_idx, val in enumerate(row_data):
            cell = row.cells[cell_idx]
            cell.width = d_widths[cell_idx]
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F9FAFC")
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_after = Pt(0)
            r = p_c.add_run(val)
            r.font.size = Pt(8)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    h6 = doc.add_heading("6. Guía de Evidencias Visuales y Capturas del Tablero de Miro", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)
    
    screens_t3 = [
        ("Captura 1: Tablero Miro - Marco 1 (Diagnóstico del Pipeline y Semáforo de 3 Columnas)", "Muestra el flujo horizontal completo, las columnas Verde, Amarillo, Rojo y las 4 preguntas clave."),
        ("Captura 2: Tablero Miro - Marco 2 (Robustez del CI y Pruebas Unitarias)", "Muestra el Mínimo Viable de Calidad, los tests agregados y el acuerdo del equipo sobre CI en rojo."),
        ("Captura 3: Tablero Miro - Marco 3 (Auto-Deploy a Render y Decisiones)", "Muestra el checklist de Render, comandos de build/start, health check path y política de despliegue."),
        ("Captura 4: Tablero Miro - Marco 4 (Feature Toggle con ConfigCat)", "Muestra la integración del SDK, el toggle wishlist_enabled, demostración OFF vs ON y gobernanza."),
        ("Captura 5: Tablero Miro - Marco 5 (Definition of Done Final en 3 Columnas)", "Muestra la tabla de DoD Técnico, DoD de Negocio y DoD de Despliegue con el cierre de Fase 1 en verde."),
        ("Captura 6: Tablero Miro - Marco 6 (Cierre del Taller y Acciones de Seguimiento)", "Muestra la ronda final, compromisos individuales y las 3 acciones de seguimiento acordadas."),
        ("Captura 7: Aplicación en Producción en Render con Estilos y ConfigCat", "Captura de la app en vivo en https://proyecto-mongodb-ventas.onrender.com mostrando los estilos y los productos cargados.")
    ]
    
    for s_title, s_desc in screens_t3:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.space_before = Pt(6)
        p_st.paragraph_format.space_after = Pt(2)
        r_st = p_st.add_run(f"📸 {s_title}")
        r_st.bold = True
        r_st.font.size = Pt(10)
        r_st.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        
        p_sd = doc.add_paragraph(s_desc, style='List Bullet')
        p_sd.paragraph_format.space_after = Pt(4)
        
        box_tbl = doc.add_table(rows=1, cols=1)
        box_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_cell = box_tbl.cell(0, 0)
        box_cell.width = Inches(6.5)
        set_cell_background(box_cell, "F2F5F9")
        set_cell_margins(box_cell, top=140, bottom=140, left=180, right=180)
        tcPr = box_cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="dashed" w:sz="8" w:space="0" w:color="A6B9D0"/><w:top w:val="dashed" w:sz="8" w:space="0" w:color="A6B9D0"/><w:right w:val="dashed" w:sz="8" w:space="0" w:color="A6B9D0"/><w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="A6B9D0"/></w:tcBorders>')
        tcPr.append(borders)
        p_box = box_cell.paragraphs[0]
        p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_box.paragraph_format.space_after = Pt(0)
        r_box = p_box.add_run(f"[ Pegar aquí la imagen de evidencia: {s_title} ]")
        r_box.font.name = "Arial"
        r_box.font.size = Pt(8.5)
        r_box.font.italic = True
        r_box.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    out_path = os.path.join("docs", "INFORME_TALLER_3_RENDER_CONFIGCAT_DOD.docx")
    doc.save(out_path)
    print("Documento Word de Taller 3 generado correctamente en:", out_path)

if __name__ == "__main__":
    build_docx_taller3()
